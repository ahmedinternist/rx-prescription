"""Drug database imported from CSV/Excel and queried through SQLite.

The doctor can import/replace the database from any CSV or Excel source (e.g.
exported from another system, a national formulary). Each row needs at least
one medicine name; everything else is optional and used to enrich the
autocomplete + prescription.

Expected columns (case-insensitive, order-independent).  At least one name is
required per row:
    generic_name   - scientific / generic name (optional when brand exists)
    brand_name     - trade / brand name (optional when scientific exists)
    strength       - e.g. "500 mg" (optional)
    form           - e.g. "tablet", "syrup", "capsule" (optional)
    category       - existing local category/favorite tag (optional)
    therapeutic_group - major therapeutic group used by the dashboard (optional)
    detailed_class - detailed drug class within that group (optional)
    notes          - free text (optional)

The source file remains portable and editable.  A disposable, indexed SQLite
cache beside it provides fast runtime searches without retaining every row in
memory.  The cache is rebuilt automatically whenever the source changes.
"""
from __future__ import annotations

import csv
import sqlite3
import unicodedata
import xlrd
from contextlib import contextmanager
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, Iterator, List, Optional, Sequence, Union

COLUMN_ALIASES = {
    "generic_name": ["generic_name", "generic", "generic name", "scientific_name",
                     "scientific name", "scientific", "inn", "name", "drug_name", "drug name",
                     "المادة", "اسم المادة", "اسم الدواء", "الاسم العلمي"],
    "brand_name": ["brand_name", "brand", "trade_name", "trade name", "trade", "الاسم التجاري"],
    "strength": ["strength", "dose", "dosage_strength"],
    "form": ["form", "dosage_form", "drug_form", "route"],
    "category": ["category", "class", "group", "atc"],
    "therapeutic_group": ["therapeutic_group", "therapeutic group", "major_class", "major class",
                            "major classification"],
    "detailed_class": ["detailed_class", "detailed class", "drug_class", "drug class", "sub_class",
                       "sub classification", "sub-classification"],
    "class_mappings": ["class_mappings", "class mappings", "multiple_classes",
                       "multiple classes"],
    "mapping_status": ["mapping_status", "mapping status", "classification_status",
                       "classification status", "confidence"],
    "notes": ["notes", "comment", "remarks"],
}


@dataclass
class Drug:
    generic_name: str
    brand_name: str = ""
    strength: str = ""
    form: str = ""
    category: str = ""
    therapeutic_group: str = ""
    detailed_class: str = ""
    class_mappings: str = ""
    mapping_status: str = ""
    notes: str = ""

    def to_dict(self) -> Dict[str, str]:
        return {
            "generic_name": self.generic_name,
            "brand_name": self.brand_name,
            "strength": self.strength,
            "form": self.form,
            "category": self.category,
            "therapeutic_group": self.therapeutic_group,
            "detailed_class": self.detailed_class,
            "class_mappings": self.class_mappings,
            "mapping_status": self.mapping_status,
            "notes": self.notes,
        }


DRUG_FIELDS = (
    "generic_name", "brand_name", "strength", "form", "category",
    "therapeutic_group", "detailed_class", "class_mappings",
    "mapping_status", "notes",
)


def _norm(text: str) -> str:
    """Lower-case + strip accents for tolerant matching."""
    text = text.lower().strip()
    return "".join(
        c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c)
    )


def _header_key(value: object) -> str:
    """Normalise spreadsheet headers without altering the displayed values."""
    return " ".join(str(value).replace("\ufeff", "").replace("\xa0", " ")
                    .replace("_", " ").casefold().split())


def _resolve_columns(header: List[str]) -> Dict[str, str]:
    """Map our canonical field names to the actual column names."""
    lower = {_header_key(h): h for h in header if str(h).strip()}
    mapping: Dict[str, str] = {}
    # Pharmacy workbooks sometimes contain both "Generic Name" (used as the
    # product/trade field) and "Scientific Name" (the INN). Prefer the explicit
    # scientific column whenever both are present.
    scientific_keys = (
        "scientific name", "scientific", "inn", "الاسم العلمي",
    )
    scientific_column = next(
        (lower[key] for key in scientific_keys if key in lower), None)
    if scientific_column:
        mapping["generic_name"] = scientific_column
    for field_name, aliases in COLUMN_ALIASES.items():
        if field_name in mapping:
            continue
        for alias in aliases:
            if _header_key(alias) in lower:
                mapping[field_name] = lower[_header_key(alias)]
                break
    if scientific_column and "brand_name" not in mapping:
        product_keys = (
            "generic / trade name", "generic/trade name",
            "generic or trade name", "generic trade name",
            "generic name", "generic",
        )
        product_column = next(
            (lower[key] for key in product_keys
             if key in lower and lower[key] != scientific_column), None)
        if product_column:
            mapping["brand_name"] = product_column
    return mapping


# ---------------------------------------------------------------------------
# File reading helpers (module level) — support CSV and Excel.
# ---------------------------------------------------------------------------
def _new_import_stats() -> Dict[str, int]:
    return {
        "source_rows": 0,
        "imported_rows": 0,
        "brand_only_rows": 0,
        "scientific_only_rows": 0,
        "skipped_blank_names": 0,
    }


def _read_rows_with_stats(path: Path) -> tuple[List[Dict[str, str]], Dict[str, int]]:
    """Read a supported database and retain row-level import diagnostics."""
    stats = _new_import_stats()
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        rows = _read_xlsx(path, stats)
    elif suffix == ".xls":
        rows = _read_xls(path, stats)
    else:
        rows = _read_csv(path, stats)
    stats["imported_rows"] = len(rows)
    return rows, stats


def _read_rows(path: Path) -> List[Dict[str, str]]:
    """Read any supported file (CSV / XLSX) into canonical-field dicts."""
    return _read_rows_with_stats(path)[0]


def _read_csv(path: Path, stats: Optional[Dict[str, int]] = None) -> List[Dict[str, str]]:
    with open(path, "r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        rows = list(reader)
    if not rows:
        return []
    header = rows[0]
    return _rows_to_dicts(header, rows[1:], stats)


def _read_xlsx(path: Path, stats: Optional[Dict[str, int]] = None) -> List[Dict[str, str]]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    grids: List[List[List[str]]] = []
    try:
        # Workbooks received from pharmacies often contain a cover sheet or a
        # report title above the column headers.  Use the first worksheet that
        # contains recognisable medicine-column headers, not just wb.active.
        for ws in wb.worksheets:
            grid = [["" if c is None else str(c) for c in row]
                    for row in ws.iter_rows(values_only=True)]
            grids.append(grid)
    finally:
        wb.close()
    return _read_grids(grids, stats)


def _read_xls(path: Path, stats: Optional[Dict[str, int]] = None) -> List[Dict[str, str]]:
    """Read legacy Excel workbooks using xlrd (which supports .xls BIFF files)."""
    workbook = xlrd.open_workbook(path)
    grids = []
    for sheet in workbook.sheets():
        grids.append([["" if value is None else str(value) for value in sheet.row_values(index)]
                      for index in range(sheet.nrows)])
    return _read_grids(grids, stats)


def _read_grids(grids: List[List[List[str]]],
                stats: Optional[Dict[str, int]] = None) -> List[Dict[str, str]]:
    """Find a medicine header row across spreadsheet sheets and title rows."""
    fallback_grid: List[List[str]] = []
    for grid in grids:
        grid = [row for row in grid if any(str(cell).strip() for cell in row)]
        if not grid:
            continue
        if not fallback_grid:
            fallback_grid = grid
        for header_index, header in enumerate(grid):
            if "generic_name" in _resolve_columns(header):
                return _rows_to_dicts(header, grid[header_index + 1:], stats)
    # Preserve support for a simple one-column sheet whose first column is the
    # medicine name, even when it does not use a recognised header label.
    if not fallback_grid:
        return []
    return _rows_to_dicts(fallback_grid[0], fallback_grid[1:], stats)


def _rows_to_dicts(header: List[str], rows: List[List[str]],
                   stats: Optional[Dict[str, int]] = None) -> List[Dict[str, str]]:
    mapping = _resolve_columns(header)
    if "generic_name" not in mapping:
        mapping["generic_name"] = header[0]
    out = []
    for row in rows:
        if not row or not any(c.strip() for c in row):
            continue
        if stats is not None:
            stats["source_rows"] += 1
        rec = {k: (row[header.index(v)].strip() if v in header and header.index(v) < len(row) else "")
               for k, v in mapping.items()}
        scientific = rec.get("generic_name", "").strip()
        brand = rec.get("brand_name", "").strip()
        if scientific or brand:
            if stats is not None:
                if brand and not scientific:
                    stats["brand_only_rows"] += 1
                elif scientific and not brand:
                    stats["scientific_only_rows"] += 1
            out.append(rec)
        elif stats is not None:
            stats["skipped_blank_names"] += 1
    return out


class DrugCollection(Sequence[Drug]):
    """Lazy, list-compatible view over the SQLite drug table."""

    def __init__(self, database: "DrugDatabase") -> None:
        self.database = database

    def __len__(self) -> int:
        return self.database.count()

    def __iter__(self) -> Iterator[Drug]:
        yield from self.database.iter_drugs()

    def __getitem__(self, index: Union[int, slice]):
        if isinstance(index, slice):
            start, stop, step = index.indices(len(self))
            if step != 1:
                return list(self)[index]
            return self.database.fetch_page(start, max(0, stop - start))
        if index < 0:
            index += len(self)
        rows = self.database.fetch_page(index, 1)
        if not rows:
            raise IndexError(index)
        return rows[0]


class DrugDatabase:
    def __init__(self, path: str) -> None:
        self.path = Path(path)
        self.cache_path = self.path.with_suffix(self.path.suffix + ".sqlite3")
        self.drugs: Sequence[Drug] = DrugCollection(self)
        self.columns: List[str] = []
        self.last_import_report: Dict[str, int] = _new_import_stats()
        self.load()

    @contextmanager
    def _connect(self):
        connection = sqlite3.connect(self.cache_path, timeout=20)
        connection.row_factory = sqlite3.Row
        try:
            yield connection
        finally:
            connection.close()

    def _source_signature(self) -> str:
        if not self.path.exists():
            return "missing"
        stat = self.path.stat()
        return f"{self.path.resolve()}|{stat.st_size}|{stat.st_mtime_ns}"

    def _cache_is_current(self) -> bool:
        if not self.cache_path.exists():
            return False
        try:
            with self._connect() as connection:
                value = connection.execute(
                    "SELECT value FROM meta WHERE key='source_signature'"
                ).fetchone()
                return bool(value and value[0] == self._source_signature())
        except sqlite3.Error:
            return False

    def _rebuild_cache(self, drugs: Sequence[Drug]) -> None:
        self.cache_path.parent.mkdir(parents=True, exist_ok=True)
        temporary = self.cache_path.with_suffix(self.cache_path.suffix + ".tmp")
        if temporary.exists():
            temporary.unlink()
        connection = sqlite3.connect(temporary)
        try:
            connection.executescript("""
                PRAGMA journal_mode=OFF;
                PRAGMA synchronous=OFF;
                CREATE TABLE meta (key TEXT PRIMARY KEY, value TEXT NOT NULL);
                CREATE TABLE drugs (
                    id INTEGER PRIMARY KEY,
                    generic_name TEXT NOT NULL,
                    brand_name TEXT NOT NULL DEFAULT '',
                    strength TEXT NOT NULL DEFAULT '',
                    form TEXT NOT NULL DEFAULT '',
                    category TEXT NOT NULL DEFAULT '',
                    therapeutic_group TEXT NOT NULL DEFAULT '',
                    detailed_class TEXT NOT NULL DEFAULT '',
                    class_mappings TEXT NOT NULL DEFAULT '',
                    mapping_status TEXT NOT NULL DEFAULT '',
                    notes TEXT NOT NULL DEFAULT '',
                    generic_norm TEXT NOT NULL,
                    brand_norm TEXT NOT NULL,
                    prescribable_norm TEXT NOT NULL,
                    search_norm TEXT NOT NULL
                );
                CREATE INDEX idx_drugs_generic ON drugs(generic_norm);
                CREATE INDEX idx_drugs_brand ON drugs(brand_norm);
                CREATE INDEX idx_drugs_prescribable ON drugs(prescribable_norm);
                CREATE INDEX idx_drugs_group ON drugs(therapeutic_group);
                CREATE INDEX idx_drugs_detail ON drugs(detailed_class);
                CREATE INDEX idx_drugs_mapping_status ON drugs(mapping_status);
            """)
            values = []
            for drug in drugs:
                record = drug.to_dict()
                generic = _norm(drug.generic_name)
                brand = _norm(drug.brand_name)
                prescribable = brand or generic
                searchable = " ".join(filter(None, (generic, brand, _norm(drug.strength),
                                                      _norm(drug.form), _norm(drug.category))))
                values.append(tuple(record[field] for field in DRUG_FIELDS)
                              + (generic, brand, prescribable, searchable))
            connection.executemany(
                "INSERT INTO drugs (" + ",".join(DRUG_FIELDS) +
                ",generic_norm,brand_norm,prescribable_norm,search_norm) "
                "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)", values)
            connection.execute(
                "INSERT INTO meta(key,value) VALUES('source_signature',?)",
                (self._source_signature(),))
            connection.commit()
        finally:
            connection.close()
        # SQLite's backup API can refresh an active cache safely even when a
        # short-lived reader is finishing on Windows (where replacing an open
        # database file would otherwise raise PermissionError).
        source = sqlite3.connect(temporary)
        destination = sqlite3.connect(self.cache_path, timeout=20)
        try:
            source.backup(destination)
        finally:
            destination.close()
            source.close()
        temporary.unlink(missing_ok=True)

    @staticmethod
    def _row_to_drug(row) -> Drug:
        return Drug(**{field: row[field] for field in DRUG_FIELDS})

    # -- loading ------------------------------------------------------------
    def load(self) -> int:
        """Open or rebuild the indexed runtime cache. Returns number loaded."""
        if not self.path.exists():
            self._rebuild_cache([])
            return 0
        self.columns = list(COLUMN_ALIASES.keys())
        if not self._cache_is_current():
            rows = _read_rows(self.path)
            loaded = [self._dict_to_drug(rec) for rec in rows
                      if (rec.get("generic_name", "").strip()
                          or rec.get("brand_name", "").strip())]
            self._rebuild_cache(loaded)
        self.drugs = DrugCollection(self)
        return self.count()

    def count(self) -> int:
        if not self.cache_path.exists():
            return 0
        with self._connect() as connection:
            return int(connection.execute("SELECT COUNT(*) FROM drugs").fetchone()[0])

    def iter_drugs(self, batch_size: int = 500) -> Iterator[Drug]:
        if not self.cache_path.exists():
            return
        with self._connect() as connection:
            cursor = connection.execute(
                "SELECT " + ",".join(DRUG_FIELDS) + " FROM drugs ORDER BY id")
            while True:
                rows = cursor.fetchmany(batch_size)
                if not rows:
                    break
                for row in rows:
                    yield self._row_to_drug(row)

    def fetch_page(self, offset: int, limit: int) -> List[Drug]:
        if limit <= 0 or offset < 0 or not self.cache_path.exists():
            return []
        with self._connect() as connection:
            rows = connection.execute(
                "SELECT " + ",".join(DRUG_FIELDS) +
                " FROM drugs ORDER BY id LIMIT ? OFFSET ?", (limit, offset)).fetchall()
        return [self._row_to_drug(row) for row in rows]

    # -- import / replace ---------------------------------------------------
    def import_file(self, source_path: str, replace: bool = True) -> int:
        """Import a CSV or Excel (XLS/XLSX) file into the active database.

        replace=True overwrites the current DB; replace=False merges
        (skipping exact duplicate product identities).
        """
        src = Path(source_path)
        if not src.exists():
            raise FileNotFoundError(f"Source file not found: {source_path}")

        rows, import_stats = _read_rows_with_stats(src)
        if not rows:
            raise ValueError("No medicine rows were found in the selected file.")

        if replace:
            self.path.parent.mkdir(parents=True, exist_ok=True)
            imported = [self._dict_to_drug(r) for r in rows]
            import_stats.update({
                "confirmed_rows": sum(
                    drug.mapping_status == "confirmed" for drug in imported),
                "suggested_rows": sum(
                    drug.mapping_status == "suggested" for drug in imported),
                "unrecognized_class_rows": sum(
                    drug.mapping_status == "unrecognized" for drug in imported),
                "unclassified_rows": sum(
                    not drug.therapeutic_group or not drug.detailed_class
                    or drug.mapping_status == "unrecognized"
                    for drug in imported),
            })
            self.last_import_report = import_stats
            self._write(imported)
            # Re-read the active local CSV so the autocomplete and class pages
            # immediately use exactly what was saved.
            return self.load()

        # Merge by the complete product identity.  This preserves separate
        # trade products and does not collapse every brand-only row into the
        # same empty-scientific-name key.
        existing = list(self.drugs)
        existing_identities = {self.drug_identity(d) for d in existing}
        added = 0
        for r in rows:
            d = self._dict_to_drug(r)
            identity = self.drug_identity(d)
            if identity not in existing_identities:
                existing.append(d)
                existing_identities.add(identity)
                added += 1
        import_stats["added_rows"] = added
        self.last_import_report = import_stats
        self._write(existing)
        self.load()
        return added

    def clear(self) -> None:
        """Remove all medicines from the active local database safely."""
        self._write([])
        self.load()

    def add_confirmed_medicine(self, name: str, therapeutic_group: str,
                               detailed_class: str) -> tuple[Drug, bool]:
        """Add a named medicine, or confirm/map an existing exact name."""
        import drug_classes as classes

        medicine_name = str(name).strip()
        group = classes.resolve_group_code(therapeutic_group)
        detail = classes.resolve_detailed_class(group, detailed_class)
        if not medicine_name:
            raise ValueError("Enter a medicine name.")
        if not group or not detail:
            raise ValueError("Choose a valid major group and detailed drug class.")

        drugs = list(self.drugs)
        existing = next(
            (drug for drug in drugs if _norm(drug.generic_name) == _norm(medicine_name)
             or _norm(drug.brand_name) == _norm(medicine_name)), None)
        created = existing is None
        medicine = existing or Drug(generic_name=medicine_name)
        if created:
            drugs.append(medicine)
        self._set_classification(medicine, group, detail, append=True)
        self._write(drugs)
        self.load()
        refreshed = next(
            (drug for drug in self.drugs
             if self.drug_identity(drug) == self.drug_identity(medicine)), medicine)
        return refreshed, created

    def delete_drug(self, target: Drug) -> bool:
        """Delete exactly one matching product row from the local database."""
        wanted = self.drug_identity(target)
        remaining: List[Drug] = []
        deleted = False
        for drug in self.drugs:
            if not deleted and self.drug_identity(drug) == wanted:
                deleted = True
                continue
            remaining.append(drug)
        if not deleted:
            return False
        self._write(remaining)
        self.load()
        return True

    @staticmethod
    def _dict_to_drug(rec: Dict[str, str]) -> "Drug":
        drug = Drug(
            generic_name=rec.get("generic_name", "").strip(),
            brand_name=rec.get("brand_name", "").strip(),
            strength=rec.get("strength", "").strip(),
            form=rec.get("form", "").strip(),
            category=rec.get("category", "").strip(),
            therapeutic_group=rec.get("therapeutic_group", "").strip(),
            detailed_class=rec.get("detailed_class", "").strip(),
            class_mappings=rec.get("class_mappings", "").strip(),
            mapping_status=rec.get("mapping_status", "").strip(),
            notes=rec.get("notes", "").strip(),
        )
        DrugDatabase._normalise_imported_classification(drug)
        return drug

    @staticmethod
    def _normalise_imported_classification(drug: "Drug") -> None:
        """Canonicalise imported mappings and reject phantom major groups."""
        import drug_classes as classes

        pairs: List[tuple[str, str]] = []
        unrecognized: List[tuple[str, str]] = []
        inferred = False

        def add_pair(raw_group: str, raw_detail: str) -> bool:
            nonlocal inferred
            legacy = classes.classify(
                drug.generic_name, drug.category, raw_group, raw_detail)
            if raw_group.strip() == "pediatric_fluids" and legacy:
                pair = (legacy.code, legacy.detail)
                if pair not in pairs:
                    pairs.append(pair)
                    inferred = True
                return True
            group = classes.resolve_group_code(raw_group)
            detail = classes.resolve_detailed_class(group, raw_detail)
            if not group and raw_detail:
                group, detail = classes.group_for_detailed_class(raw_detail)
                inferred = inferred or bool(group)
            # Keep explicitly supplied but unknown class text visible for
            # clinician review.  It must not silently become "Other".
            if group and raw_detail.strip() and not detail:
                unrecognized.append((group, raw_detail.strip()))
                return False
            if group and not detail:
                detail = classes.resolve_detailed_class(group, drug.category)
            if group and not detail:
                suggestion = (classes.classify(drug.generic_name, drug.category)
                              or classes.classify(drug.brand_name, drug.category))
                if suggestion and suggestion.code == group:
                    detail = suggestion.detail
                    inferred = True
            if group and not detail:
                configured = classes.subclasses_for(group)
                if len(configured) == 1:
                    detail = configured[0]
                    inferred = True
                elif "Other" in configured:
                    detail = "Other"
                    inferred = True
            pair = (group, detail)
            if all(pair) and pair not in pairs:
                pairs.append(pair)
                return True
            return False

        for item in str(drug.class_mappings or "").split("|"):
            group, separator, detail = item.partition("::")
            if separator:
                add_pair(group.strip(), detail.strip())
        if drug.therapeutic_group or drug.detailed_class:
            add_pair(drug.therapeutic_group, drug.detailed_class)

        if not pairs and unrecognized:
            drug.therapeutic_group, drug.detailed_class = unrecognized[0]
            drug.class_mappings = ""
            drug.mapping_status = "unrecognized"
            return

        if not pairs:
            suggestion = (classes.classify(drug.generic_name, drug.category)
                          or classes.classify(drug.brand_name, drug.category))
            if suggestion:
                pairs.append((suggestion.code, suggestion.detail))
                inferred = True

        if not pairs:
            drug.therapeutic_group = ""
            drug.detailed_class = ""
            drug.class_mappings = ""
            drug.mapping_status = ""
            return

        drug.therapeutic_group, drug.detailed_class = pairs[0]
        drug.class_mappings = "|".join(
            f"{group}::{detail}" for group, detail in pairs)
        status = str(drug.mapping_status or "").casefold()
        drug.mapping_status = (
            "suggested" if inferred else
            status if status in {"confirmed", "suggested"} else "confirmed"
        )

    def export_csv(self, dest_path: str) -> None:
        """Export the current database to a CSV file."""
        self._write(self.drugs, Path(dest_path))

    def export_file(self, dest_path: str) -> None:
        """Export the database as CSV, modern XLSX, or legacy XLS."""
        path = Path(dest_path)
        suffix = path.suffix.casefold()
        fields = list(DRUG_FIELDS)
        path.parent.mkdir(parents=True, exist_ok=True)
        if suffix == ".csv":
            self._write(self.drugs, path)
            return
        if suffix == ".xlsx":
            from openpyxl import Workbook
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Drug Database"
            sheet.append(fields)
            for drug in self.drugs:
                record = drug.to_dict()
                sheet.append([record.get(field, "") for field in fields])
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            workbook.save(path)
            return
        if suffix == ".xls":
            import xlwt
            workbook = xlwt.Workbook(encoding="utf-8")
            sheet = workbook.add_sheet("Drug Database")
            for column, field in enumerate(fields):
                sheet.write(0, column, field)
            for row, drug in enumerate(self.drugs, 1):
                record = drug.to_dict()
                for column, field in enumerate(fields):
                    sheet.write(row, column, record.get(field, ""))
            workbook.save(str(path))
            return
        raise ValueError("Choose CSV (.csv), Excel (.xlsx), or legacy Excel (.xls).")

    def _write(self, drugs: List[Drug], path: Optional[Path] = None) -> None:
        path = path or self.path
        path.parent.mkdir(parents=True, exist_ok=True)
        fields = list(DRUG_FIELDS)
        temporary = path.with_suffix(path.suffix + ".tmp")
        with open(temporary, "w", encoding="utf-8", newline="") as fh:
            writer = csv.DictWriter(fh, fieldnames=fields)
            writer.writeheader()
            for d in drugs:
                writer.writerow(d.to_dict())
        temporary.replace(path)

    # -- search / autocomplete ---------------------------------------------
    def search(self, query: str, limit: int = 25) -> List[Drug]:
        """Return drugs matching a query (prefix first, then substring)."""
        q = _norm(query)
        if not q:
            return self.fetch_page(0, limit)
        return self._search_column("search_norm", q, limit)

    def search_scientific(self, query: str, limit: int = 25) -> List[Drug]:
        """Find medicines by the scientific / INN name only."""
        return self._search_field(query, "generic_name", limit)

    def search_trade(self, query: str, limit: int = 25) -> List[Drug]:
        """Find medicines by the generic / trade product name only."""
        return self._search_field(query, "brand_name", limit)

    def search_prescribable(self, query: str, limit: int = 25) -> List[Drug]:
        """Search the name users actually select in the prescription form.

        Modern two-column databases use ``brand_name`` for that field. Older
        and current one-column imports store the product name in
        ``generic_name`` and leave ``brand_name`` empty, so those rows are
        intentionally searched by their only available name.
        """
        return self._search_column("prescribable_norm", _norm(query), limit)

    def _search_field(self, query: str, field_name: str, limit: int) -> List[Drug]:
        column = "generic_norm" if field_name == "generic_name" else "brand_norm"
        return self._search_column(column, _norm(query), limit)

    def _search_column(self, column: str, query: str, limit: int) -> List[Drug]:
        if not query or limit <= 0:
            return []
        if column not in {"generic_norm", "brand_norm", "prescribable_norm", "search_norm"}:
            raise ValueError("Unsupported search column")
        upper = query + "\uffff"
        fields = ",".join(DRUG_FIELDS)
        with self._connect() as connection:
            rows = connection.execute(
                f"SELECT {fields} FROM drugs WHERE {column}>=? AND {column}<? "
                f"ORDER BY {column}, id LIMIT ?", (query, upper, limit)).fetchall()
            remaining = limit - len(rows)
            if remaining:
                rows += connection.execute(
                    f"SELECT {fields} FROM drugs WHERE instr({column},?)>0 "
                    f"AND NOT ({column}>=? AND {column}<?) "
                    f"ORDER BY {column}, id LIMIT ?",
                    (query, query, upper, remaining)).fetchall()
        return [self._row_to_drug(row) for row in rows]

    def trade_names_for_scientific(self, scientific_name: str) -> List[str]:
        """Return the locally recorded trade names for one scientific name."""
        target = _norm(scientific_name)
        with self._connect() as connection:
            rows = connection.execute(
                "SELECT DISTINCT brand_name FROM drugs WHERE generic_norm=? AND brand_name<>'' "
                "ORDER BY brand_norm", (target,)).fetchall()
        names = {row[0].strip() for row in rows if row[0].strip()}
        return sorted(names, key=str.casefold)

    def all_names(self) -> List[str]:
        with self._connect() as connection:
            return [row[0] for row in connection.execute(
                "SELECT generic_name FROM drugs ORDER BY id")]

    def find_exact(self, name: str) -> Optional[Drug]:
        """Return a database medicine when its generic name is an exact match."""
        target = _norm(name)
        with self._connect() as connection:
            row = connection.execute(
                "SELECT " + ",".join(DRUG_FIELDS) +
                " FROM drugs WHERE generic_norm=? ORDER BY id LIMIT 1", (target,)).fetchone()
        return self._row_to_drug(row) if row else None

    def contains_name(self, generic_name: str = "", brand_name: str = "") -> bool:
        generic, brand = _norm(generic_name), _norm(brand_name)
        if not generic and not brand:
            return False
        clauses, values = [], []
        if generic:
            clauses.append("generic_norm=?")
            values.append(generic)
        if brand:
            clauses.append("brand_norm=?")
            values.append(brand)
        with self._connect() as connection:
            return connection.execute(
                "SELECT 1 FROM drugs WHERE " + " OR ".join(clauses) + " LIMIT 1",
                values).fetchone() is not None

    @staticmethod
    def drug_identity(drug: Drug) -> tuple[str, str, str, str]:
        """Return the stable identity of one imported product row."""
        return tuple(_norm(getattr(drug, field, "")) for field in (
            "generic_name", "brand_name", "strength", "form"))

    def update_classification(self, name: str, therapeutic_group: str,
                              detailed_class: str, append: bool = True) -> bool:
        """Persist a clinician-reviewed class mapping for one local medicine."""
        drugs = list(self.drugs)
        drug = next((item for item in drugs if _norm(item.generic_name) == _norm(name)), None)
        if drug is None:
            return False
        self._set_classification(drug, therapeutic_group, detailed_class, append)
        self._write(drugs)
        self.load()
        return True

    def update_classifications(self, names, therapeutic_group: str,
                               detailed_class: str, append: bool = True) -> int:
        """Persist one reviewed classification for several selected medicines."""
        wanted = {_norm(str(name)) for name in names if str(name).strip()}
        if not wanted:
            return 0
        updated = 0
        drugs = list(self.drugs)
        for drug in drugs:
            if _norm(drug.generic_name) not in wanted:
                continue
            self._set_classification(drug, therapeutic_group, detailed_class, append)
            updated += 1
        if updated:
            self._write(drugs)
            self.load()
        return updated

    def update_drug_classifications(self, selected_drugs, therapeutic_group: str,
                                    detailed_class: str, append: bool = True) -> int:
        """Persist mappings for the exact selected product rows only."""
        wanted = {self.drug_identity(drug) for drug in selected_drugs}
        if not wanted:
            return 0
        updated = 0
        drugs = list(self.drugs)
        for drug in drugs:
            if self.drug_identity(drug) not in wanted:
                continue
            self._set_classification(drug, therapeutic_group, detailed_class, append)
            updated += 1
        if updated:
            self._write(drugs)
            self.load()
        return updated

    def classification_states(self, names) -> list[Dict[str, str]]:
        """Capture classification fields so a mapping change can be recovered."""
        wanted = {_norm(str(name)) for name in names if str(name).strip()}
        return [{key: value for key, value in drug.to_dict().items()
                 if key in {"generic_name", "therapeutic_group", "detailed_class",
                            "class_mappings", "mapping_status"}}
                for drug in self.drugs if _norm(drug.generic_name) in wanted]

    def classification_states_for_drugs(self, selected_drugs) -> list[Dict[str, str]]:
        """Capture recoverable mapping state for exact selected product rows."""
        wanted = {self.drug_identity(drug) for drug in selected_drugs}
        identity_fields = {"generic_name", "brand_name", "strength", "form"}
        mapping_fields = {"therapeutic_group", "detailed_class",
                          "class_mappings", "mapping_status"}
        return [{key: value for key, value in drug.to_dict().items()
                 if key in identity_fields | mapping_fields}
                for drug in self.drugs if self.drug_identity(drug) in wanted]

    def restore_classification_states(self, states) -> int:
        valid_states = [item for item in states
                        if (isinstance(item, dict)
                            and (item.get("generic_name") or item.get("brand_name")))]
        exact = {
            tuple(_norm(str(item.get(field, ""))) for field in (
                "generic_name", "brand_name", "strength", "form")): item
            for item in valid_states
            if any(str(item.get(field, "")).strip()
                   for field in ("brand_name", "strength", "form"))
        }
        legacy = {_norm(str(item.get("generic_name", ""))): item
                  for item in valid_states
                  if not any(str(item.get(field, "")).strip()
                             for field in ("brand_name", "strength", "form"))}
        if not exact and not legacy:
            return 0
        drugs = list(self.drugs)
        updated = 0
        for drug in drugs:
            state = exact.get(self.drug_identity(drug)) or legacy.get(_norm(drug.generic_name))
            if not state:
                continue
            for key in ("therapeutic_group", "detailed_class", "class_mappings", "mapping_status"):
                setattr(drug, key, str(state.get(key, "")))
            updated += 1
        if updated:
            self._write(drugs)
            self.load()
        return updated

    @staticmethod
    def _stored_class_pairs(drug: Drug) -> List[tuple[str, str]]:
        pairs: List[tuple[str, str]] = []
        for item in str(drug.class_mappings or "").split("|"):
            group, separator, detail = item.partition("::")
            pair = (group.strip(), detail.strip())
            if separator and all(pair) and pair not in pairs:
                pairs.append(pair)
        legacy = (drug.therapeutic_group.strip(), drug.detailed_class.strip())
        if all(legacy) and legacy not in pairs:
            pairs.insert(0, legacy)
        return pairs

    @classmethod
    def _set_classification(cls, drug: Drug, therapeutic_group: str,
                            detailed_class: str, append: bool) -> None:
        pair = (therapeutic_group.strip(), detailed_class.strip())
        pairs = (cls._stored_class_pairs(drug)
                 if append and drug.mapping_status != "unrecognized" else [])
        if all(pair) and pair not in pairs:
            pairs.append(pair)
        if pairs:
            drug.therapeutic_group, drug.detailed_class = pairs[0]
        else:
            drug.therapeutic_group, drug.detailed_class = pair
        drug.class_mappings = "|".join(f"{group}::{detail}" for group, detail in pairs)
        drug.mapping_status = "confirmed"
