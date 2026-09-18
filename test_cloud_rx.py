"""Cloud QR tests: mocked networking and isolated Windows application data."""
import io
import json
from urllib.error import HTTPError, URLError

import pytest

import cloud_rx as cloud
from test_core import run_isolated


class Response(io.BytesIO):
    status = 201


def test_cloud_post_utf8_headers_and_exact_returned_link():
    payload = {"date": "2026-09-18", "doctor": "د. أحمد", "patient": "أحمد علي",
               "medications": [{"tradeName": "Brand", "instructions": "مع الطعام"}]}
    url = "https://rx-v2.vercel.app/p/fictitious-test"
    calls = []
    def opener(request, timeout):
        calls.append(request)
        assert request.full_url == cloud.API_URL and request.method == "POST"
        assert timeout == 15
        assert request.get_header("Content-type") == "application/json"
        assert request.get_header("X-api-key") == "test-key"
        assert json.loads(request.data.decode("utf-8")) == payload
        assert "أحمد".encode("utf-8") in request.data
        return Response(json.dumps({"rxId": "test", "url": url}).encode())
    result = cloud.upload_prescription(payload, " test-key ", opener=opener)
    assert result == cloud.CloudRxLink("test", url) and len(calls) == 1


@pytest.mark.parametrize("status,code", [(401, "auth"), (403, "auth"), (429, "quota"),
                                        (500, "server"), (503, "server"), (400, "http"), (302, "redirect")])
def test_cloud_http_errors_do_not_disclose_server_body(status, code):
    def opener(*args, **kwargs):
        raise HTTPError(cloud.API_URL, status, "PRIVATE SERVER DETAIL", {}, io.BytesIO(b"patient or key"))
    with pytest.raises(cloud.CloudRxError) as caught:
        cloud.upload_prescription({}, "private-key", opener=opener)
    assert caught.value.code == code and caught.value.status == status
    assert "PRIVATE" not in str(caught.value) and "private-key" not in str(caught.value)


@pytest.mark.parametrize("error,code", [(TimeoutError(), "timeout"), (URLError(TimeoutError()), "timeout"),
                                       (URLError("private detail"), "offline"), (OSError(), "offline")])
def test_cloud_connection_errors(error, code):
    def opener(*args, **kwargs):
        raise error
    with pytest.raises(cloud.CloudRxError) as caught:
        cloud.upload_prescription({}, "test", opener=opener)
    assert caught.value.code == code


@pytest.mark.parametrize("result", [[], {}, {"rxId": "", "url": "https://rx-v2.vercel.app/p/a"},
    {"rxId": "a", "url": "http://rx-v2.vercel.app/p/a"},
    {"rxId": "a", "url": "https://evil.example/p/a"},
    {"rxId": "a", "url": "https://rx-v2.vercel.app/p/"},
    {"rxId": "a", "url": "https://rx-v2.vercel.app/p/a?data=secret"},
    {"rxId": "a", "url": "https://rx-v2.vercel.app/p/a#secret"},
    {"rxId": "a", "url": "https://user@rx-v2.vercel.app/p/a"},
    {"rxId": "a", "url": "https://rx-v2.vercel.app:444/p/a"}])
def test_cloud_invalid_response(result):
    with pytest.raises(cloud.CloudRxError) as caught:
        cloud.upload_prescription({}, "test", opener=lambda *a, **k: Response(json.dumps(result).encode()))
    assert caught.value.code == "response"


@pytest.mark.parametrize("raw", [b"not-json", b"\xff", b"x" * 65537],
                         ids=["invalid-json", "invalid-utf8", "oversized"])
def test_cloud_unreadable_response(raw):
    with pytest.raises(cloud.CloudRxError) as caught:
        cloud.upload_prescription({}, "test", opener=lambda *a, **k: Response(raw))
    assert caught.value.code == "response"


def test_cloud_missing_key_and_redirects_are_not_sent():
    with pytest.raises(cloud.CloudRxError) as caught:
        cloud.upload_prescription({}, "", opener=lambda *a, **k: pytest.fail("must not send"))
    assert caught.value.code == "missing_key"
    assert cloud._NoRedirect().redirect_request(None, None, 302, "", {}, "https://evil.example") is None


def test_cloud_minimal_payload_and_encrypted_credentials(tmp_path):
    run_isolated(tmp_path, '''
from config import config, Config, CONFIG_PATH
from qr_utils import Prescription, Doctor, Patient, Clinic, DrugItem
rx = Prescription(clinic=Clinic(name="Private clinic",logo_path="private-path"),
    doctor=Doctor(name="Doctor",license_no="private-license",specialty="GP"),
    patient=Patient(name="Patient",age="50",sex="F",allergies="private"),
    drugs=[DrugItem(brand_name="Brand only",notes="مع الطعام")],date="2026-09-18")
payload = rx.to_qr_payload()
assert payload["v"] == 4 and payload["doctor"] == {"name":"Doctor","specialty":"GP"}
assert payload["patient"] == {"name":"Patient"}
assert payload["drugs"] == [{"brand_name":"Brand only","notes":"مع الطعام"}]
assert "clinic" not in payload and "sig" not in payload
config.data["signing_private_key"] = "old-key-preserved"
config.cloud_rx_api_key = "fictitious-encrypted-test-key"
raw = CONFIG_PATH.read_text()
assert "fictitious-encrypted-test-key" not in raw and '"dpapi-v1"' in raw
reloaded = Config()
assert reloaded.cloud_rx_api_key == "fictitious-encrypted-test-key"
assert reloaded.data["signing_private_key"] == "old-key-preserved"
''')


def test_cloud_failure_choices_and_closure(tmp_path):
    run_isolated(tmp_path, '''
from types import SimpleNamespace
from main import App
calls=[]
ui=SimpleNamespace(_closing=False,
    _upload_export_link=lambda op:calls.append(("retry",op)),
    _write_cloud_export=lambda op,url:calls.append(("without",op,url)),
    _set_export_busy=lambda value:calls.append(("busy",value)))
op=object()
for choice in ("retry","without","cancel"):
    App._resolve_cloud_export_failure(ui,op,choice)
assert calls == [("retry",op),("without",op,None),("busy",False)]
ui._closing=True
App._resolve_cloud_export_failure(ui,op,"retry")
assert len(calls)==3
''')


def test_mobile_viewer_payload_exact_fields_and_privacy(tmp_path):
    run_isolated(tmp_path, '''
from qr_utils import Prescription, Doctor, Patient, Clinic, DrugItem
rx = Prescription(clinic=Clinic(name="Do not upload",address="Do not upload",phone=" 07700000000 ",logo_path="private-path"),
    doctor=Doctor(name=" د. أحمد ",license_no=" TEST ",specialty=" باطنية "),
    patient=Patient(name=" أحمد علي ",age=" 50 ",sex="F",id_number="private-id",allergies="private"),
    drugs=[DrugItem(brand_name=" Brand ",generic_name=" Molecule ",dosage=" 5 mg ",
                   frequency="1x1",duration="7 days",notes="مع الطعام",quantity="7"),
           DrugItem(brand_name="Brand only"),DrugItem(generic_name="Scientific only")],
    date="2026-09-18",rx_id="private-local-id",diagnosis="private",refills="1")
payload = rx.to_cloud_payload()
assert payload == {"doctor":"د. أحمد · باطنية","registrationId":"TEST","phone":"07700000000",
    "patient":"أحمد علي","age":"50","date":"2026-09-18","medications":[
    {"tradeName":"Brand","genericName":"Molecule","dosage":"5 mg","instructions":"1x1 · مع الطعام",
     "duration":"7 days","quantity":"7"},{"tradeName":"Brand only"},{"genericName":"Scientific only"}]}
rx.doctor.license_no = rx.clinic.phone = rx.patient.age = " "
rx.drugs = [DrugItem(brand_name="Brand",notes="عند النوم"),DrugItem(generic_name="Molecule",frequency="1x2")]
payload = rx.to_cloud_payload()
assert set(payload) == {"doctor","patient","date","medications"}
assert payload["medications"] == [{"tradeName":"Brand","instructions":"عند النوم"},
                                 {"genericName":"Molecule","instructions":"1x2"}]
rx.patient.age = "0"
assert rx.to_cloud_payload()["age"] == "0"
''')


def test_cloud_export_snapshot_exact_qr_and_busy_guard(tmp_path):
    run_isolated(tmp_path, '''
from types import SimpleNamespace
import main, i18n as I
from main import App
from config import config
from qr_utils import Prescription, Doctor, Patient, DrugItem
rx=Prescription(doctor=Doctor(name="Original",license_no="1"),patient=Patient(name="Original patient"),
    drugs=[DrugItem(brand_name="Original brand")],date="2026-09-18")
config.cloud_rx_api_key="test-key"
config.data["document_defaults"]={"language":"interface","margin_mm":16}
I.set_lang("en")
calls=[]
ui=SimpleNamespace(_export_busy=False,_closing=False,collect=lambda:rx,_validate=lambda *a:True,
    paper_var=SimpleNamespace(get=lambda:"A5"),_upload_export_link=lambda op:calls.append(op))
ui._prepare_full_document=lambda action:App._prepare_full_document(ui,action)
ui._set_export_busy=lambda value:setattr(ui,"_export_busy",value)
op=App._start_cloud_export(ui,"Export",path_docx="captured.docx")
assert ui._export_busy and len(calls)==1
assert App._start_cloud_export(ui,"Export") is None
rx.patient.name="Different patient"
rx.drugs[0].brand_name="Different brand"
config.data["document_defaults"]["margin_mm"]=24
I.set_lang("ar")
assert op["payload"]["patient"]=="Original patient"
assert op["payload"]["medications"]==[{"tradeName":"Original brand"}]
assert op["rx"].drugs[0].brand_name=="Original brand"
assert op["document"]["margin_mm"]==16 and op["document"]["language"]=="en"
assert op["document"]["_paper_size"]=="A5"
import threading
ui._document_lock=threading.Lock()
seen=[]
main.qu.make_qr_image=lambda url:seen.append(url) or "QR image"
ui._generate_full_document=lambda rx,qr,doc,*paths:seen.append((qr,doc,paths))
ui.submit_background=lambda worker,finished,**kw:finished(worker())
op["on_success"]=lambda result:seen.append("finished")
App._write_cloud_export(ui,op,"https://rx-v2.vercel.app/p/exact")
assert seen[0]=="https://rx-v2.vercel.app/p/exact" and seen[1][0]=="QR image"
assert not ui._export_busy and seen[-1]=="finished"
seen.clear()
App._write_cloud_export(ui,op,None)
assert isinstance(seen[0],tuple) and seen[0][0] is None
''')


def test_cloud_document_language_does_not_change_interface():
    import concurrent.futures
    import i18n as I
    I.set_lang("en")
    def worker():
        with I.document_language("ar"):
            assert I.get_lang() == "ar"
            return I.t("cloud_cancel")
    with concurrent.futures.ThreadPoolExecutor() as executor:
        assert executor.submit(worker).result() == "إلغاء"
    assert I.get_lang() == "en"


def test_background_completion_uses_ui_queue_and_ignores_closed_app(tmp_path):
    run_isolated(tmp_path, '''
import queue, threading, concurrent.futures
from types import SimpleNamespace
from main import App
executor=concurrent.futures.ThreadPoolExecutor(max_workers=1)
ui=SimpleNamespace(_closing=False,_background_callbacks=queue.SimpleQueue(),_executor=executor)
seen=[]
main_thread=threading.get_ident()
future=App.submit_background(ui,lambda:threading.get_ident(),lambda result:seen.append((result,threading.get_ident())),silent=True)
future.result(timeout=3)
executor.shutdown(wait=True)
assert not seen
ui._background_callbacks.get(timeout=2)()
assert seen[0][0]!=main_thread and seen[0][1]==main_thread
executor=concurrent.futures.ThreadPoolExecutor(max_workers=1)
ui._executor=executor
future=App.submit_background(ui,lambda:"result",lambda result:seen.append("must not run"),silent=True)
future.result(timeout=3)
executor.shutdown(wait=True)
ui._closing=True
ui._background_callbacks.get(timeout=2)()
assert len(seen)==1
''')


def test_cloud_a4_a5_exports_with_and_without_qr(tmp_path):
    run_isolated(tmp_path, '''
import re, zipfile, os
from pathlib import Path
from docx import Document
import pdf_generator as pdf, qr_utils as q, i18n as I
I.set_lang("en")
root=Path(os.environ["RX_APP_DATA_DIR"])
rx=q.Prescription(doctor=q.Doctor(name="Fictitious doctor",license_no="TEST"),
    patient=q.Patient(name="Fictitious patient"),drugs=[q.DrugItem(brand_name="Fictitious brand",frequency="1x1")])
qr=q.make_qr_image("https://rx-v2.vercel.app/p/fictitious-test")
for paper,points in (("A5",(419.528,595.276)),("A4",(595.276,841.889))):
    for include in (False,True):
        image=qr if include else None
        for kind,generator in (("header",pdf.generate_prescription_docx),("compact",pdf.generate_medication_label_docx)):
            path=root/f"{paper}-{kind}-{include}.docx"
            generator(rx,path,paper_size=paper,qr_pil_image=image)
            doc=Document(path)
            assert abs(doc.sections[0].page_width.pt-points[0])<.1
            assert abs(doc.sections[0].page_height.pt-points[1])<.1
            assert len(doc.inline_shapes)==int(include)
            assert any(I.t("pdf_scan") in p.text for p in doc.paragraphs)==include
            if not include and kind=="compact":
                assert doc.paragraphs[-1].text.startswith("1.")
        path=root/f"{paper}-{include}.pdf"
        pdf.generate_prescription_pdf(rx,path,paper_size=paper,qr_pil_image=image)
        raw=path.read_bytes()
        media=re.search(rb"/MediaBox\\s*\\[\\s*0\\s+0\\s+([\\d.]+)\\s+([\\d.]+)",raw)
        assert media and abs(float(media[1])-points[0])<.1 and abs(float(media[2])-points[1])<.1
        assert (b"/Subtype /Image" in raw)==include
''')


def test_cloud_actual_ui_async_states_settings_and_close(tmp_path):
    run_isolated(tmp_path, '''
import time, threading
import main, cloud_rx, i18n as I
from main import App, SettingsWindow, VisualButton
from config import config
from qr_utils import Prescription, Doctor, DrugItem
config.cloud_rx_api_key="fictitious-test-key"
app=App()
app.update()
errors=[]
app.report_callback_exception=lambda *args:errors.append(args)
app.collect=lambda:Prescription(doctor=Doctor(name="Fictitious",license_no="TEST"),drugs=[DrugItem(brand_name="Test")])
app._generate_full_document=lambda *args:None
gate=threading.Event()
ticks=[]
def upload(*args):
    assert gate.wait(3)
    return cloud_rx.CloudRxLink("test","https://rx-v2.vercel.app/p/fictitious")
main.cloud_rx.upload_prescription=upload
op=app._start_cloud_export("Export",path_docx="test.docx")
assert app._export_busy
assert all(w.cget("state")=="disabled" for w in app.action.winfo_children() if isinstance(w,VisualButton))
app.after(15,lambda:ticks.append("responsive"))
app.after(80,gate.set)
deadline=time.monotonic()+4
while app._export_busy and time.monotonic()<deadline:
    app.update()
    time.sleep(.005)
assert not app._export_busy and ticks and not errors
assert all(w.cget("state")=="normal" for w in app.action.winfo_children() if isinstance(w,VisualButton))
settings=SettingsWindow(app)
app.update()
assert settings.cloud_key_entry.cget("show")=="•"
assert not hasattr(settings,"viewer_var")
settings.destroy()
app._set_export_busy(True)
app._cloud_export_failed(op,cloud_rx.CloudRxError("offline"))
app.update()
dialog=next(w for w in app.winfo_children() if isinstance(w,main.ctk.CTkToplevel))
pending=[dialog]
buttons=[]
while pending:
    w=pending.pop()
    if isinstance(w,VisualButton):buttons.append(w)
    pending.extend(w.winfo_children())
assert {b.cget("text") for b in buttons}=={I.t("cloud_retry"),I.t("cloud_without_qr"),I.t("cloud_cancel")}
next(b for b in buttons if b.cget("text")==I.t("cloud_cancel")).invoke()
assert not app._export_busy
gate.clear()
app._start_cloud_export("Export",path_docx="never-created.docx")
app.destroy()
gate.set()
app._executor.shutdown(wait=True)
assert app._closing and not errors
''')
