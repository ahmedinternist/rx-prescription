from pathlib import Path

from test_core import run_isolated


def test_layout_profiles_are_normalized_and_printer_specific():
    import print_layout

    profiles = print_layout.normalize_profiles({
        "Clinic printer": {
            "horizontal_offset_mm": 99,
            "qr_position": "Right",
            "qr_size_mm": 28,
        }
    })
    assert print_layout.DEFAULT_PROFILE_NAME in profiles
    assert profiles["Clinic printer"]["horizontal_offset_mm"] == 25
    document = {
        "selected_printer": "Clinic printer",
        "printer_profiles": profiles,
    }
    selected = print_layout.active_profile(document)
    assert selected["qr_position"] == "Right"
    assert selected["qr_size_mm"] == 28


def test_named_calibration_presets_do_not_inject_printer_default():
    import print_layout

    presets = print_layout.normalize_presets({
        "Clinic A5": {"qr_position": "Right", "medication_top_mm": 88},
        "": {"qr_position": "Center"},
    })
    assert list(presets) == ["Clinic A5"]
    assert print_layout.DEFAULT_PROFILE_NAME not in presets
    assert presets["Clinic A5"]["qr_position"] == "Right"
    assert presets["Clinic A5"]["medication_top_mm"] == 88


def test_named_calibration_presets_persist_in_encrypted_settings(tmp_path: Path):
    run_isolated(tmp_path, '''
import importlib
import config

config.config.data["document_defaults"]["calibration_presets"] = {
    "Clinic paper": {"medication_top_mm": 91, "qr_position": "Center"}}
config.config.save()
reloaded = config.Config()
presets = reloaded.data["document_defaults"]["calibration_presets"]
assert presets["Clinic paper"]["medication_top_mm"] == 91
assert presets["Clinic paper"]["qr_position"] == "Center"
assert config.DEFAULT_PROFILE_NAME not in presets
''')


def test_custom_a5_alignment_controls_medication_and_qr(tmp_path: Path):
    run_isolated(tmp_path, '''
import os, zipfile
from pathlib import Path
from docx import Document
import pdf_generator as pdf
import qr_utils as q

root=Path(os.environ["RX_APP_DATA_DIR"])
path=root/"custom-layout.docx"
rx=q.Prescription(drugs=[q.DrugItem(
    brand_name="Medicine", dosage="10 mg", frequency="1x1")])
layout={
    "horizontal_offset_mm":2,
    "vertical_offset_mm":3,
    "medication_top_mm":90,
    "line_spacing":1.5,
    "field_gap_mm":12,
    "qr_position":"Right",
    "qr_size_mm":30,
    "qr_side_margin_mm":10,
    "qr_bottom_margin_mm":25,
}
pdf.generate_medication_label_docx(
    rx,path,paper_size="A5",layout_profile=layout,
    qr_pil_image=q.make_qr_image("https://rx-v2.vercel.app/p/test1234"))
doc=Document(path)
med=next(p for p in doc.paragraphs if p.text.startswith("1."))
assert 80.5 < med.paragraph_format.space_before.mm < 81.5
assert abs(float(med.paragraph_format.line_spacing)-1.5)<.01
assert "Medicine" in med.text and "10 mg" in med.text and "1x1" in med.text
with zipfile.ZipFile(path) as archive:
    xml=archive.read("word/document.xml").decode("utf-8")
assert '<wp:positionH relativeFrom="page"><wp:posOffset>3960000</wp:posOffset>' in xml
assert '<wp:positionV relativeFrom="page"><wp:posOffset>5688000</wp:posOffset>' in xml
''')


def test_calibration_page_contains_no_patient_data(tmp_path: Path):
    run_isolated(tmp_path, '''
import os, zipfile
from pathlib import Path
from docx import Document
import pdf_generator as pdf

root=Path(os.environ["RX_APP_DATA_DIR"])
path=root/"calibration.docx"
pdf.generate_calibration_docx(path,paper_size="A5",layout_profile={
    "medication_top_mm":86,"qr_position":"Left","qr_size_mm":32})
doc=Document(path)
assert abs(doc.sections[0].page_width.mm-148)<.2
assert abs(doc.sections[0].page_height.mm-210)<.2
with zipfile.ZipFile(path) as archive:
    xml=archive.read("word/document.xml").decode("utf-8")
assert "patient" not in xml.casefold()
assert "doctor" not in xml.casefold()
assert "<wp:anchor" in xml
''')


def test_word_filename_uses_patient_and_current_date():
    import datetime
    from types import SimpleNamespace
    from main import App

    ui = SimpleNamespace(patient_vars={
        "name": SimpleNamespace(get=lambda: "Ahmed Ali")})
    assert App._word_export_filename(ui) == (
        f"Ahmed Ali_{datetime.date.today().isoformat()}.docx")
